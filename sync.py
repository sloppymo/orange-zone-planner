#!/usr/bin/env python3
"""Parse 'The Orange Zone Sheet' xlsx export into state.js for the planner.

Usage:
    python3 sync.py [path/to/export.xlsx]

With no argument, picks the newest 'The Orange Zone Sheet*.xlsx' in ~/Downloads
by modification time. Writes state.js next to this script.
"""
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
OUT = HERE / "state.js"

SKILL_ROWS = {
    "Strength": "STR", "Dexterity": "DEX", "Constitution": "CON", "Agility": "AGI",
    "Crafting": "CRF", "Fieldwork": "FLD", "Stealth": "STL", "Scavenging": "SCAV",
    "Intelligence": "INT", "Willpower": "WILL", "Charisma": "CHA", "Perception": "PER",
    "Physics": "PHYS", "Biology": "BIO", "Medicine": "MED", "Computers": "COMP",
}
STATUS_ROWS = ["Injury", "Plague", "Stress", "Hunger", "Fatigue"]
NUM_CHARS = 6  # columns B..G

# keyword (lowercase, in the 'N slot(s)' header line) -> station id
STATION_KEYS = {
    "the shelter": "shelter",
    "shelter": "shelter",
    "generator bike": "bike",
    "water tanks": "water",
    "the forge": "forge",
    "forge": "forge",
    "the workshop": "workshop",
    "workshop": "workshop",
    "sigint": "sigint",
    "hydroponics": "hydro",
    "logging": "logging",
    "iron processing": "iron",
    "lookout station": "lookout",
    "lookout": "lookout",
}
EXPECTED_STATIONS = {"shelter", "bike", "water", "forge", "workshop", "sigint", "hydro", "logging", "iron", "lookout"}
HEADER_RE = re.compile(r"-\s*(?:\d+|unlimited)\s+slots?\b", re.I)


def num(v, default=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def split_trait(text):
    if not text or not str(text).strip():
        return None
    s = str(text).strip()
    m = re.match(r"([A-Za-z0-9'\" -]+?)\s*(?:[-–—:]\s+)(.*)", s, re.S)
    if m:
        return {"name": m.group(1).strip(), "text": m.group(2).strip()}
    return {"name": s, "text": s}


def norm_hydration(val):
    if not val:
        return "FULL"
    tok = str(val).strip().upper().split()[0]
    return tok if tok in {"FULL", "LIGHT", "MODERATE", "SEVERE"} else "FULL"


def pick_source():
    files = list(Path.home().joinpath("Downloads").glob("The Orange Zone Sheet*.xlsx"))
    if not files:
        sys.exit("No 'The Orange Zone Sheet*.xlsx' in ~/Downloads; pass a path")
    return str(max(files, key=lambda p: p.stat().st_mtime))


def parse_stations(wb):
    """Pull condition / abundance / queue / reveal progress from Workstations tab."""
    ws = wb["Workstations"]
    st = {}
    cur = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if not v:
            continue
        t = str(v).strip()
        tl = t.lower()
        if HEADER_RE.search(tl):
            cur = None
            for k, sid in sorted(STATION_KEYS.items(), key=lambda kv: -len(kv[0])):
                if k in tl:
                    cur = sid
                    st.setdefault(sid, {})
                    break
            if not cur:
                print(f"WARNING: unrecognized station header at Workstations!A{r}: {t!r}", file=sys.stderr)
                continue
            m = re.search(r"\((\d+)\s*(?:-\s*Repair|\))", t, re.I)
            if m:
                st[cur]["cond"] = int(m.group(1))
            continue
        if not cur:
            continue
        m = re.search(r"Condition:?\s*(\d+)", t, re.I)
        if m:
            st[cur]["cond"] = int(m.group(1))
        m = re.search(r"Abundance:?\s*(\d+)", t, re.I)
        if m:
            st[cur]["abundance"] = int(m.group(1))
        m = re.search(r"reveal progress:?\s*\((\d+)/(\d+)\)", t, re.I)
        if m:
            st[cur]["progress"], st[cur]["need"] = int(m.group(1)), int(m.group(2))
        m = re.search(r"[Qq]ueue:\s*(.+)", t)
        if m:
            st[cur]["queue"] = m.group(1).strip()
            mb = re.search(r"\((\d+)/(\d+)(?:\s*banked)?\)", t, re.I)
            if mb:
                st[cur]["banked"], st[cur]["need"] = int(mb.group(1)), int(mb.group(2))
        else:
            mb = re.search(r"\((\d+)/(\d+)\s*banked\)", t, re.I)
            if mb:
                st[cur]["banked"], st[cur]["need"] = int(mb.group(1)), int(mb.group(2))
    missing = EXPECTED_STATIONS - set(st)
    if missing:
        print(f"WARNING: missing stations from Workstations tab: {sorted(missing)}", file=sys.stderr)
    for sid, d in st.items():
        if "cond" not in d:
            print(f"WARNING: station '{sid}' has no parsed condition; defaulting to 100", file=sys.stderr)
    return st


def parse_day(wb):
    for name in wb.sheetnames:
        if name.strip().lower() == "day":
            v = wb[name].cell(row=1, column=2).value
            if v is not None:
                return num(v, 1.0)
    print("WARNING: no 'Day' tab found; defaulting day to 1", file=sys.stderr)
    return 1.0


def parse_weather(wb):
    ws = wb["Weather"]
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and "current weather" in str(v).strip().lower():
            nxt = ws.cell(row=r + 1, column=1).value
            if nxt:
                return str(nxt).strip()
    v = ws.cell(row=29, column=1).value
    return str(v).strip() if v else "Sunny"


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else pick_source()
    print(f"Reading {src}")
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["Character Sheets"]

    labels = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is not None:
            labels[str(v).strip()] = r

    trait_start = labels.get("TRAITS", 47)
    cond_start = labels.get("CONDITIONS", 51)
    grade_start = next((r for k, r in labels.items() if k.upper().startswith("WORKSTATION GRADES")), 58)

    chars = []
    for i in range(NUM_CHARS):
        col = 2 + i
        name = ws.cell(row=labels["Name"], column=col).value
        if not name:
            continue
        skills = {}
        for label, code in SKILL_ROWS.items():
            raw = ws.cell(row=labels[label], column=col).value
            if isinstance(raw, str) and "/" in raw:
                cur, mx = raw.split("/", 1)
            else:
                cur = mx = raw
            skills[code] = {"cur": num(cur), "max": num(mx)}
        statuses = {s.lower(): num(ws.cell(row=labels[s], column=col).value) for s in STATUS_ROWS}
        traits = [t for r in range(trait_start, cond_start)
                  if (t := split_trait(ws.cell(row=r, column=col).value))]
        conditions = [str(ws.cell(row=r, column=col).value).strip()
                      for r in range(cond_start, grade_start)
                      if ws.cell(row=r, column=col).value and str(ws.cell(row=r, column=col).value).strip()]
        chars.append({
            "name": str(name).strip(),
            "archetype": str(ws.cell(row=labels["Archetype"], column=col).value or "").strip(),
            "xp": str(ws.cell(row=labels["XP to next Level"], column=col).value or "").strip(),
            "statuses": statuses,
            "hydration": norm_hydration(ws.cell(row=labels["HYDRATION"], column=col).value),
            "skills": skills,
            "traits": traits,
            "conditions": conditions,
        })

    grade_rows = {}
    for r in range(grade_start, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if label:
            grade_rows[str(label).strip()] = [
                str(ws.cell(row=r, column=2 + i).value or "--").strip() for i in range(len(chars))
            ]

    inv = {}
    wsi = wb["Inventory"]
    for r in range(1, wsi.max_row + 1):
        a, b = wsi.cell(row=r, column=1).value, wsi.cell(row=r, column=2).value
        if a and b is not None:
            inv[str(a).strip()] = b.strip() if isinstance(b, str) else b

    state = {
        "source": Path(src).name,
        "day": parse_day(wb),
        "weather": parse_weather(wb),
        "power": {"current": num(inv["Power"], 3.0) if "Power" in inv else 3.0, "max": 5, "base": 3},
        "inventory": inv,
        "stations": parse_stations(wb),
        "characters": chars,
        "sheetGrades": grade_rows,
    }
    state["hash"] = hashlib.sha256(json.dumps(state, sort_keys=True).encode()).hexdigest()[:12]
    OUT.write_text("// Generated by sync.py from %s — do not edit by hand\nconst SHEET_STATE = %s;\n"
                   % (state["source"], json.dumps(state, indent=1)))
    print(f"Wrote {OUT} ({len(chars)} characters, day {state['day']}, weather {state['weather']}, "
          f"hash {state['hash']}, stations: {sorted(state['stations'])})")


if __name__ == "__main__":
    main()
